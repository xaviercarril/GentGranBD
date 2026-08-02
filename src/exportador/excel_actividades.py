from __future__ import annotations

from math import ceil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HEADERS = [
    "Activitat",
    "Professor/a",
    "Preu matrícula",
    "Inscrits",
    "Màx. alumnes",
    "Descripció",
]
COLUMN_WIDTHS = [34, 30, 16, 12, 14, 55]

TITLE_FILL = PatternFill("solid", fgColor="548235")
HEADER_FILL = PatternFill("solid", fgColor="D9E6C5")
ALTERNATE_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN_GRAY = Side(style="thin", color="D9D9D9")


def _money_value(value):
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace("€", "").strip().replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return str(value)


def _integer_value(value):
    if value in (None, ""):
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return value


def _row_height(values: list) -> float:
    text_columns = {0: 42, 1: 38, 5: 75}
    lines = 1
    for index, chars_per_line in text_columns.items():
        text = str(values[index] or "")
        wrapped_lines = sum(
            max(1, ceil(len(line) / chars_per_line))
            for line in text.splitlines() or [""]
        )
        lines = max(lines, wrapped_lines)
    return max(24, (lines * 15) + 4)


def generar_excel_actividades_curso(
    curso_nombre: str,
    actividades: list[dict],
    ruta_xlsx: str,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Cursos"
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A4"
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.print_title_rows = "1:3"

    title = f"LLISTAT DE CURSOS - {curso_nombre}" if curso_nombre else "LLISTAT DE CURSOS"
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    title_cell = worksheet.cell(row=1, column=1, value=title)
    title_cell.fill = TITLE_FILL
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 26

    for column, header in enumerate(HEADERS, start=1):
        cell = worksheet.cell(row=3, column=column, value=header)
        cell.fill = HEADER_FILL
        cell.font = Font(name="Calibri", size=11, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN_GRAY, right=THIN_GRAY)
    worksheet.row_dimensions[3].height = 28

    for row_number, actividad in enumerate(actividades, start=4):
        values = [
            actividad.get("nombre", "") or "",
            actividad.get("personal_nombre", "") or "",
            _money_value(actividad.get("precio_matricula")),
            _integer_value(actividad.get("inscritos")),
            _integer_value(actividad.get("numMaxAlumnos")),
            actividad.get("descripcion", "") or "",
        ]
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(row=row_number, column=column, value=value)
            cell.alignment = Alignment(
                horizontal="right" if column in {3, 4, 5} else "left",
                vertical="top",
                wrap_text=column in {1, 2, 6},
                indent=1 if column in {1, 2, 6} else 0,
            )
            cell.border = Border(bottom=THIN_GRAY, right=THIN_GRAY)
            if row_number % 2 == 1:
                cell.fill = ALTERNATE_FILL
        worksheet.cell(row=row_number, column=3).number_format = '#,##0.00 "€"'
        worksheet.cell(row=row_number, column=4).number_format = "#,##0"
        worksheet.cell(row=row_number, column=5).number_format = "#,##0"
        worksheet.row_dimensions[row_number].height = _row_height(values)

    for column, width in enumerate(COLUMN_WIDTHS, start=1):
        worksheet.column_dimensions[get_column_letter(column)].width = width

    last_row = max(3, len(actividades) + 3)
    worksheet.auto_filter.ref = f"A3:F{last_row}"
    worksheet.print_area = f"A1:F{last_row}"

    output = Path(ruta_xlsx)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
