from datetime import date

from openpyxl import load_workbook
from sqlalchemy import create_engine, inspect, text
from sqlalchemy.orm import sessionmaker

import controladores.actividades as actividades
import controladores.inscripcion_socio as inscripciones
import database
from exportador.excel_participantes import generar_excel_participantes_viaje_session
from exportador.pdf_inscripciones import TRIP_HEADERS, _build_table
from models import Actividad, Base, EstadoInscripcion, InscripcionSocio, TipoActividadEnum
from ui.actividad_detail import InscripcionesActividadTableModel


def test_campos_del_viaje_se_guardan_y_se_pueden_modificar(monkeypatch):
    engine = create_engine("sqlite:///:memory:", echo=False)
    Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine, expire_on_commit=False)
    monkeypatch.setattr(actividades, "SessionLocal", Session)
    monkeypatch.setattr(inscripciones, "SessionLocal", Session)

    with Session() as session:
        viaje = Actividad(nombre="Viatge", tipo=TipoActividadEnum.VIATGE)
        session.add(viaje)
        session.commit()
        viaje_id = viaje.id

    inscripcion = inscripciones.registrar_inscripcion(
        {
            "actividadID": viaje_id,
            "noSocioNombre": "Maria",
            "fechaInscripcion": date.today(),
            "estado": EstadoInscripcion.INSCRIT.value,
            "asiento": "12A",
            "lugarRecogida": "Plaça de l'Església",
        }
    )

    guardada = inscripciones.consultar_inscripcion(inscripcion.id)
    assert guardada["asiento"] == "12A"
    assert guardada["lugarRecogida"] == "Plaça de l'Església"

    inscripciones.modificar_inscripcion(
        inscripcion.id,
        {"asiento": "8", "lugarRecogida": "Ajuntament"},
    )

    modificada = inscripciones.consultar_inscripcion(inscripcion.id)
    assert modificada["asiento"] == "8"
    assert modificada["lugarRecogida"] == "Ajuntament"
    engine.dispose()


def test_campos_del_viaje_son_editables_en_la_tabla():
    assert {"asiento", "lugarRecogida"} <= InscripcionesActividadTableModel.EDITABLE_KEYS


def test_excel_del_viaje_incluye_asiento_y_lugar_de_recogida(session, tmp_path):
    viaje = Actividad(nombre="Viatge", tipo=TipoActividadEnum.VIATGE)
    session.add(viaje)
    session.commit()
    session.add(
        InscripcionSocio(
            actividadID=viaje.id,
            noSocioNombre="Joan",
            fechaInscripcion=date.today(),
            estado=EstadoInscripcion.INSCRIT,
            asiento="5B",
            lugarRecogida="Estació",
        )
    )
    session.commit()

    output = tmp_path / "participants.xlsx"
    generar_excel_participantes_viaje_session(session, viaje.id, str(output))

    ws = load_workbook(output)["Participants"]
    headers = [cell.value for cell in ws[3]]
    assert headers[9:11] == ["Seient", "Lloc de recollida"]
    assert ws["J4"].value == "5B"
    assert ws["K4"].value == "Estació"


def test_pdf_del_viaje_incluye_asiento_y_lugar_de_recogida():
    inscripcion = InscripcionSocio(
        fechaInscripcion=date.today(),
        estado=EstadoInscripcion.INSCRIT,
        asiento="3C",
        lugarRecogida="Centre cívic",
    )

    table = _build_table([inscripcion], include_dni=True)
    row = [cell.getPlainText() for cell in table._cellvalues[1]]

    assert row[TRIP_HEADERS.index("Seient")] == "3C"
    assert row[TRIP_HEADERS.index("Lloc de recollida")] == "Centre cívic"


def test_migracion_anade_campos_del_viaje_a_bases_existentes(monkeypatch):
    engine = create_engine("sqlite:///:memory:", echo=False)
    with engine.begin() as connection:
        connection.execute(text("CREATE TABLE socios (id INTEGER PRIMARY KEY)"))
        connection.execute(text("CREATE TABLE inscripciones (id INTEGER PRIMARY KEY)"))

    monkeypatch.setattr(database, "engine", engine)
    database.ensure_schema_updates()

    columns = {column["name"] for column in inspect(engine).get_columns("inscripciones")}
    assert {"asiento", "lugarRecogida"} <= columns
    engine.dispose()
